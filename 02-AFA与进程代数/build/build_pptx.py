# -*- coding: utf-8 -*-
"""Build the lecture slide deck (PPTX)."""
import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image

ROOT = r"D:\AI\AFARules"
FIG = os.path.join(ROOT, "figures")
OUT = os.path.join(ROOT, "out")

NAVY = RGBColor(0x0F, 0x17, 0x2A)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
AMBER = RGBColor(0xB4, 0x53, 0x09)
LIGHT = RGBColor(0xF8, 0xFA, 0xFC)
GRAY = RGBColor(0x64, 0x74, 0x8B)
PALE = RGBColor(0xCB, 0xD5, 0xE1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

MATH_CHARS = set("ΩτωαβγδεθλμνξπρστφχψΔΣΦΨ∈∉⊆⊂∅℘→←↔⇒⇔⟨⟩≈∼≅≡≠≤≥×·∘⋃⋂∑∏∞⊤⊥")


def set_ea(run, name):
    rPr = run._r.get_or_add_rPr()
    for el in rPr.findall(qn("a:ea")):
        rPr.remove(el)
    ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", name)


def add_runs(p, text, size=18, color=NAVY, bold=False, ascii_font="Arial", ea_font="Microsoft YaHei"):
    parts = re.split(r"([\u4e00-\u9fff]+)", text)
    for part in parts:
        if not part:
            continue
        run = p.add_run()
        run.text = part
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        if re.match(r"^[\u4e00-\u9fff]+$", part):
            run.font.name = ea_font
            set_ea(run, ea_font)
        else:
            if any(ch in MATH_CHARS for ch in part):
                run.font.name = "Cambria Math"
                set_ea(run, ea_font)
            else:
                run.font.name = ascii_font
                set_ea(run, ea_font)


def textbox(slide, left, top, width, height, text, size=18, color=NAVY, bold=False,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.15):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        add_runs(p, line, size=size, color=color, bold=bold)
    return box


def bullets(slide, left, top, width, height, items, size=18, color=NAVY, space=10):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, item in enumerate(items):
        level = 0
        if isinstance(item, tuple):
            level, item = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = level
        p.space_after = Pt(space)
        p.line_spacing = 1.12
        mark = "•  " if level == 0 else "–  "
        mark_color = BLUE if level == 0 else GRAY
        r = p.add_run()
        r.text = mark
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = mark_color
        r.font.name = "Arial"
        add_runs(p, item, size=size if level == 0 else size - 2, color=color)
    return box


def bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def accent_bar(slide, left=0.9, top=0.85, width=0.09, height=0.55, color=BLUE):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top),
                                 Inches(width), Inches(height))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()


def header(slide, kicker, title):
    textbox(slide, 1.15, 0.55, 11.5, 0.35, kicker, size=12, color=GRAY, bold=True)
    textbox(slide, 1.15, 0.9, 11.5, 0.9, title, size=30, color=NAVY, bold=True)
    accent_bar(slide, left=0.85, top=0.95, height=0.7)
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1.15), Inches(1.78),
                                      Inches(12.4), Inches(1.78))
    line.line.color.rgb = PALE
    line.line.width = Pt(0.8)


def footer(slide, idx, dark=False):
    color = PALE if dark else GRAY
    textbox(slide, 1.0, 7.05, 6.0, 0.3, "AFA 与进程代数 · 2026", size=10, color=color)
    textbox(slide, 11.4, 7.05, 1.2, 0.3, f"{idx:02d}", size=10, color=color, align=PP_ALIGN.RIGHT)


def add_image_fit(slide, path, left, top, width, height):
    iw, ih = Image.open(path).size
    scale = min(width / iw, height / ih)
    w, h = iw * scale, ih * scale
    x = left + (width - w) / 2
    y = top + (height - h) / 2
    slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))


def title_slide(prs, idx):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, NAVY)
    # motif
    pts = [(10.3, 2.0), (11.55, 1.25), (12.25, 2.55)]
    for i in range(3):
        a, b = pts[i], pts[(i + 1) % 3]
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(a[0]), Inches(a[1]),
                                    Inches(b[0]), Inches(b[1]))
        ln.line.color.rgb = RGBColor(0x3B, 0x82, 0xF6)
        ln.line.width = Pt(2)
    for (x, y) in pts:
        o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - 0.14), Inches(y - 0.14),
                               Inches(0.28), Inches(0.28))
        o.fill.solid(); o.fill.fore_color.rgb = NAVY
        o.line.color.rgb = RGBColor(0x93, 0xC5, 0xFD); o.line.width = Pt(2)

    textbox(s, 1.1, 1.55, 8.8, 0.4, "A F A　与　进 程 代 数", size=14, color=RGBColor(0x60, 0xA5, 0xFA), bold=True)
    textbox(s, 1.05, 2.0, 9.5, 1.2, "AFA 与进程代数", size=48, color=WHITE, bold=True)
    textbox(s, 1.1, 3.25, 9.5, 0.7, "从反基础公理到超集合语义", size=24, color=RGBColor(0xCB, 0xD5, 0xE1))
    rule = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1.15), Inches(4.15),
                                  Inches(7.8), Inches(4.15))
    rule.line.color.rgb = BLUE; rule.line.width = Pt(4)
    textbox(s, 1.1, 4.45, 9.0, 0.5, "“每个图都有唯一的装饰。”", size=18, color=RGBColor(0xE2, 0xE8, 0xF0))
    textbox(s, 1.1, 4.9, 9.0, 0.4, "—— Peter Aczel, 1988", size=13, color=GRAY)
    textbox(s, 1.1, 6.35, 6.0, 0.4, "自学教材 · 16 周课程", size=13, color=RGBColor(0x94, 0xA3, 0xB8))
    footer(s, idx, dark=True)


def section_slide(prs, idx, num, title, sub):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, NAVY)
    textbox(s, 1.1, 2.1, 3.2, 1.6, num, size=96, color=RGBColor(0x1D, 0x4E, 0xD8), bold=True)
    textbox(s, 1.15, 3.9, 9.0, 0.9, title, size=38, color=WHITE, bold=True)
    textbox(s, 1.18, 4.85, 9.0, 0.6, sub, size=18, color=RGBColor(0x94, 0xA3, 0xB8))
    rule = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1.15), Inches(4.6),
                                  Inches(4.2), Inches(4.6))
    rule.line.color.rgb = BLUE; rule.line.width = Pt(3)
    footer(s, idx, dark=True)


def content_slide(prs, idx, kicker, title, items, size=19):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, WHITE)
    header(s, kicker, title)
    bullets(s, 1.2, 2.15, 11.6, 4.5, items, size=size)
    footer(s, idx)
    return s


def image_slide(prs, idx, kicker, title, caption, img, items=None, img_w=7.0):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, WHITE)
    header(s, kicker, title)
    if items:
        bullets(s, 1.2, 2.2, 4.6, 4.4, items, size=17)
    add_image_fit(s, os.path.join(FIG, img), 6.0 if items else 2.2, 2.0, img_w, 4.5)
    textbox(s, 1.2, 6.55, 11.2, 0.35, caption, size=12, color=GRAY)
    footer(s, idx)
    return s


def table_slide(prs, idx, kicker, title, rows):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, WHITE)
    header(s, kicker, title)
    shape = s.shapes.add_table(len(rows), 2, Inches(1.2), Inches(2.15), Inches(11.2), Inches(4.2))
    table = shape.table
    table.columns[0].width = Inches(3.4)
    table.columns[1].width = Inches(7.8)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.margin_left = Inches(0.12)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.05)
            cell.margin_bottom = Inches(0.05)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if i == 0 else (LIGHT if i % 2 else WHITE)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            add_runs(p, val, size=16, color=WHITE if i == 0 else NAVY, bold=(i == 0))
    footer(s, idx)
    return s


def quote_slide(prs, idx, kicker, quote, author):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, LIGHT)
    textbox(s, 1.2, 1.0, 4.0, 0.4, kicker, size=13, color=BLUE, bold=True)
    textbox(s, 1.2, 2.2, 10.8, 2.4, quote, size=34, color=NAVY, bold=True, line_spacing=1.2)
    textbox(s, 1.25, 4.9, 10.0, 0.5, author, size=16, color=GRAY)
    footer(s, idx)
    return s


def end_slide(prs, idx):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, NAVY)
    textbox(s, 1.1, 1.8, 8.0, 0.9, "谢谢", size=56, color=WHITE, bold=True)
    textbox(s, 1.15, 3.0, 10.5, 0.6, "问题、讨论、项目选题，欢迎随时交流。", size=20, color=RGBColor(0xCB, 0xD5, 0xE1))
    textbox(s, 1.15, 4.2, 10.5, 1.8,
            "必读：Aczel (1988) · Milner (1989) · Rutten (2000)\n"
            "工具：mCRL2 · CADP · FDR4\n"
            "进阶：Sangiorgi (2011) · Jacobs (2016)",
            size=16, color=RGBColor(0x94, 0xA3, 0xB8), line_spacing=1.6)
    footer(s, idx, dark=True)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    i = 1
    title_slide(prs, i); i += 1

    image_slide(prs, i, "课程地图", "四篇十三章，一条主线", "图 0-1　本书结构",
                "fig-map.png", None, img_w=11.0); i += 1

    section_slide(prs, i, "01", "非良基集合论", "从基础公理的困境到超全集"); i += 1

    content_slide(prs, i, "导论", "一个方程引出的世界", [
        "集合：x = {x}",
        "进程：X = a.X",
        "流：s = ⟨0, s⟩",
        "共同形式：x = f(x)",
        "ZFC：无解；AFA：恰有一解",
    ]); i += 1

    content_slide(prs, i, "第 1 章", "ZFC 的答案：无解", [
        "基础公理（FA）：∈ 是良基的，不存在无穷降链",
        "x = {x} 推出 x ∈ x，与 FA 矛盾",
        "s = ⟨0, s⟩ 推出无穷降链 s ∋ {0,s} ∋ s ∋ …",
        "但“无解”是一条公理的选择，不是逻辑必然",
    ]); i += 1

    content_slide(prs, i, "第 1 章", "两种集合观", [
        "迭代：V₀ = ∅，V(α+1) = ℘(Vα)，V = ⋃ Vα",
        "余迭代：集合由它的观察图决定",
        "一个自下而上，一个自上而下",
        "AFA 选择后者：允许循环，要求唯一",
    ]); i += 1

    content_slide(prs, i, "第 2 章", "ZFA 与 AFA", [
        "ZFA = ZFC − FA + 原子",
        "AFA：每个图都有唯一的装饰",
        "相对一致：ZFC 一致 ⟹ ZFA + AFA 一致",
        "AFA 蕴含 ¬FA（Ω ∈ Ω），但不蕴含有全集",
    ]); i += 1

    content_slide(prs, i, "第 2 章", "图与装饰", [
        "图 G = (V, E)，边 v → w",
        "装饰 d(v) = { d(w) | v → w }",
        "唯一性：图完全决定集合",
        "自环装饰为 Ω = {Ω}",
        "带出口的循环装饰为 Ω_∅ = {∅, Ω_∅}",
    ]); i += 1

    content_slide(prs, i, "第 2 章", "反基础公理", [
        "每个图都有唯一的装饰",
        "存在性：循环对象合法",
        "唯一性：循环对象确定",
        "一条一阶语句，不是公理模式",
    ]); i += 1

    image_slide(prs, i, "第 2 章", "纯循环都装饰为 Ω", "图 2-1　自环、双节点循环、三点循环",
                "fig-omega.png", [
                    "常值 Ω 装饰满足方程",
                    "由唯一性，它就是唯一的装饰",
                    "双节点循环与自环代表同一个集合",
                ], img_w=6.8); i += 1

    content_slide(prs, i, "第 3 章", "互模拟", [
        "关系 R：你走一步，我跟一步（zig-zag）",
        "最大互模拟 ∼ 是等价关系",
        "所有叶子都等价；所有纯循环都等价",
        "游戏：攻击者选一步，防守者匹配",
    ]); i += 1

    content_slide(prs, i, "第 3 章", "装饰定理", [
        "d(x) = d(y) ⟺ x ∼ y",
        "证明：ker(d) 是互模拟；互模拟 ⊆ ker(d)",
        "集合相等 = 行为等价",
        "商图 G/∼ 总是外延的",
    ]); i += 1

    content_slide(prs, i, "第 3 章", "超全集", [
        "V_A = { d(v) | v 是某个图的节点 }",
        "每个集合都是其 ∈-图的装饰",
        "V_A 是真类（Russell 悖论仍然有效）",
        "良基集合 WF ⊂ V_A",
    ]); i += 1

    content_slide(prs, i, "第 4 章", "解引理", [
        "方程系统 x ≈ a_x",
        "有限项：把方程画成图，取装饰",
        "每个系统有唯一解",
        "x = {x} → Ω；x = {∅, x} → Ω_∅",
    ]); i += 1

    content_slide(prs, i, "第 4 章", "余归纳", [
        "要证 x = y：找互模拟 R，使 x R y",
        "无需基始；闭合即可",
        "例：Ω = {Ω}",
        "余归纳对偶于归纳：一个兜底，一个盖楼",
    ]); i += 1

    content_slide(prs, i, "第 5 章", "余代数", [
        "函子 F；余代数 (X, ξ: X → F(X))",
        "同态：θ ∘ f = F(f) ∘ ξ",
        "图 = ℘-余代数；LTS = ℘(L×−)-余代数",
        "装饰 = 到终余代数的同态",
    ]); i += 1

    content_slide(prs, i, "第 5 章", "终余代数", [
        "每个余代数都唯一地指向它",
        "(V_A, id) 是 ℘ 的终余代数",
        "终性 = 装饰的存在唯一性",
        "ZFC 中没有终余代数（否则 Ω 存在）",
    ]); i += 1

    section_slide(prs, i, "02", "进程代数", "LTS · CCS · 弱互模拟 · ACP · π-演算"); i += 1

    content_slide(prs, i, "第 6 章", "进程与 LTS", [
        "标号迁移系统 (S, Act, →)",
        "迹：动作序列；迹等价太粗",
        "a.(b+c) 与 a.b+a.c 的迹相同，行为不同",
        "需要保留分支：互模拟",
    ]); i += 1

    image_slide(prs, i, "第 6 章", "强互模拟", "图 6-3　P = a.0|b.0 与 Q = a.b.0+b.a.0",
                "fig-bisim.png", [
                    "编号相同 = 互模拟对应",
                    "P 与 Q 行为一致",
                    "但 a.(b+c) 与 a.b+a.c 不互模拟",
                ], img_w=6.8); i += 1

    content_slide(prs, i, "第 6 章", "Hennessy–Milner 逻辑", [
        "φ ::= ⊤ | ¬φ | φ ∧ φ | ⟨a⟩φ",
        "像有限时：∼ = HML 等价",
        "区分公式：⟨a⟩(⟨b⟩⊤ ∧ ⟨c⟩⊤)",
        "注意 ⟨a⟩ 是“存在”，不是“所有”",
    ]); i += 1

    image_slide(prs, i, "第 6 章", "行为等价的谱系", "图 6-4　常见行为等价（粗略谱系）",
                "fig-spectrum.png", [
                    "迹 ⊂ 失败 ⊂ … ⊂ 互模拟（粗略）",
                    "实际构成偏序，不是全序",
                    "互模拟：同余、可判定、逻辑刻画",
                ], img_w=6.9); i += 1

    content_slide(prs, i, "第 7 章", "CCS 语法", [
        "P ::= 0 | a.P | P+P | P|P | P\\L | P[f] | rec X.P",
        "守护：每个 X 都在前缀之下",
        "互补动作 a 与 ā 通信产生 τ",
        "τ 是内部动作，外部不可见",
    ]); i += 1

    content_slide(prs, i, "第 7 章", "CCS 语义", [
        "a.P --a--> P",
        "P+Q：两条 Sum 规则",
        "P|Q：交错 + 通信（a 与 ā → τ）",
        "P\\L：a ∉ L 才能走",
        "rec X.P：展开后走一步",
    ]); i += 1

    content_slide(prs, i, "第 7 章", "展开律", [
        "P|Q ∼ Σ a.(P'|Q) + Σ a.(P|Q') + Σ τ.(P'|Q')",
        "a.0|b.0 ∼ a.b.0 + b.a.0",
        "并行 = 交错 + 通信",
        "对有限进程成立；是定理，不是公理",
    ]); i += 1

    content_slide(prs, i, "第 7 章", "递归与 RSP", [
        "RDP：rec X.P ∼ P[rec X.P/X]",
        "RSP：守护 + Q ∼ P[Q/X] ⟹ Q ∼ rec X.P",
        "守护排除 X = X 这类空转",
        "超集合模型验证 RSP：解唯一",
    ]); i += 1

    content_slide(prs, i, "第 8 章", "弱互模拟", [
        "τ 不可见：P ⇒ P' 是 τ 闭包",
        "P --a⇒ P'：⇒ --a--> ⇒",
        "≈：用弱迁移匹配",
        "a.0 ≈ τ.a.0",
    ]); i += 1

    content_slide(prs, i, "第 8 章", "τ 定律与发散", [
        "a.τ.P ≈ a.P",
        "P + τ.P ≈ τ.P",
        "Ω = rec X.τ.X ≈ 0",
        "发散被忽略；≈ 对 + 不同余",
        "观察同余 ≈^c 修复同余性",
    ]); i += 1

    content_slide(prs, i, "第 9 章", "ACP：代数进程论", [
        "代数先行：签名 + 公理",
        "动作集 A，通信函数 γ",
        "算子：+、·、∥、L、|、∂_H、τ_I",
        "交错语义；mCRL2 的理论基础",
    ]); i += 1

    content_slide(prs, i, "第 9 章", "ACP 核心公理", [
        "x∥y = L(x,y) + L(y,x) + x|y",
        "L(a,x) = a·x",
        "a|b = γ(a,b)",
        "∂_H 禁止动作；τ_I 隐藏动作",
        "RDP/RSP：守护递归有唯一解",
    ]); i += 1

    content_slide(prs, i, "第 10 章", "π-演算：移动性", [
        "名字可以传递：ā⟨b⟩.P 与 a(x).P",
        "(νa)P：私有名字",
        "作用域挤出：发送后作用域扩展",
        "早/晚/开互模拟：对名字实例化不同",
    ]); i += 1

    section_slide(prs, i, "03", "统一", "进程即超集合"); i += 1

    content_slide(prs, i, "第 11 章", "进程即超集合", [
        "d(s) = { (a, d(s')) | s --a--> s' }",
        "s ∼ t ⟺ d(s) = d(t)",
        "(V_A(Act), id) 是终余代数",
        "强互模拟 = 集合相等",
    ]); i += 1

    table_slide(prs, i, "第 11 章", "指称表",
                [["进程", "指称"],
                 ["0", "∅"],
                 ["a.0", "{(a, ∅)}"],
                 ["a.(b.0+c.0)", "{(a, {(b,∅), (c,∅)})}"],
                 ["a.b.0+a.c.0", "{(a,{(b,∅)}), (a,{(c,∅)})}"],
                 ["Ω_τ", "{(τ, Ω_τ)}"]]); i += 1

    image_slide(prs, i, "第 11 章", "语义三角", "图 11-1　语法、商与行为",
                "fig-triangle.png", [
                    "先商去互模拟，再取指称",
                    "等于直接取指称",
                    "指称在商上单射",
                ], img_w=5.8); i += 1

    content_slide(prs, i, "总结", "三句话", [
        "AFA：每个图都有唯一的装饰",
        "互模拟：集合相等",
        "终语义：进程 = 超集合",
        "边界：弱互模拟、发散、绑定名需要更丰富的结构",
    ]); i += 1

    end_slide(prs, i)

    os.makedirs(OUT, exist_ok=True)
    path = os.path.join(OUT, "AFA与进程代数-课程讲义.pptx")
    prs.save(path)
    print("saved", path, "slides", len(prs.slides))


if __name__ == "__main__":
    build()
