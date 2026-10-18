# -*- coding: utf-8 -*-
"""Build the textbook DOCX from the Markdown sources."""
import json
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(__file__))
from tex import math_segments, split_italic

ROOT = r"D:\AI\AFARules"
BOOK = os.path.join(ROOT, "book")

INK = RGBColor(0x0F, 0x17, 0x2A)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
AMBER = RGBColor(0xB4, 0x53, 0x09)
GRAY = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

BOX = {
    "definition": ("E8F5E9", "2E7D32"),
    "theorem": ("E3F2FD", "1565C0"),
    "lemma": ("E3F2FD", "1565C0"),
    "corollary": ("E3F2FD", "1565C0"),
    "proposition": ("E3F2FD", "1565C0"),
    "example": ("FFF8E1", "F9A825"),
    "remark": ("F5F5F5", "64748B"),
    "note": ("F5F5F5", "64748B"),
    "reading": ("F3E5F5", "6A1B9A"),
    "history": ("F3E5F5", "6A1B9A"),
    "algorithm": ("ECFEFF", "0E7490"),
    "proof": (None, "94A3B8"),
}

BODY_PT = 10.5
CONTENT_W = Cm(16.2)


# ----------------------------------------------------------------------------
# small ooxml helpers
# ----------------------------------------------------------------------------
def set_run_font(run, ascii_name, ea_name, size=None, bold=None, italic=None, color=None):
    run.font.name = ascii_name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), ascii_name)
    rFonts.set(qn("w:hAnsi"), ascii_name)
    rFonts.set(qn("w:eastAsia"), ea_name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade_paragraph(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def border_paragraph(p, color="2E7D32", sz=18, side="left", space=8):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    el = OxmlElement("w:" + side)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    pBdr.append(el)
    pPr.append(pBdr)


def cell_shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def cell_borders(cell, color="2E7D32", sz=18, left_only=True):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + edge)
        if left_only and edge != "left":
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def cell_margins(cell, top=120, bottom=120, left=200, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement("w:" + side)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trPr.append(el)


def add_field(paragraph, instr, placeholder=""):
    run = paragraph.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    run._r.append(fld)
    run2 = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    run2._r.append(it)
    run3 = paragraph.add_run()
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "separate")
    run3._r.append(fld2)
    if placeholder:
        paragraph.add_run(placeholder)
    run4 = paragraph.add_run()
    fld3 = OxmlElement("w:fldChar")
    fld3.set(qn("w:fldCharType"), "end")
    run4._r.append(fld3)


# ----------------------------------------------------------------------------
# parsing
# ----------------------------------------------------------------------------
SPECIAL_START = ("#", ":::", "$$", "![", "```", "> ", "|")


def is_cjk(ch):
    return "\u4e00" <= ch <= "\u9fff" or "\u3000" <= ch <= "\u303f"


def smart_join(lines):
    out = ""
    for line in lines:
        line = line.strip()
        if not out:
            out = line
        elif out and is_cjk(out[-1]) and line and is_cjk(line[0]):
            out += line
        else:
            out += " " + line
    return out


def parse_blocks(lines):
    blocks = []
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith("# "):
            blocks.append(("h1", s[2:].strip()))
            i += 1
        elif s.startswith("## "):
            blocks.append(("h2", s[3:].strip()))
            i += 1
        elif s.startswith("### "):
            blocks.append(("h3", s[4:].strip()))
            i += 1
        elif s.startswith(":::"):
            m = re.match(r":::\s*(\w+)\s*(.*)", s)
            kind = m.group(1) if m else "remark"
            title = m.group(2).strip() if m else ""
            i += 1
            body = []
            while i < n and not lines[i].strip().startswith(":::"):
                body.append(lines[i])
                i += 1
            i += 1
            blocks.append(("box", kind, title, parse_blocks(body)))
        elif s.startswith("$$"):
            if s.endswith("$$") and len(s) > 4:
                blocks.append(("math", s[2:-2].strip()))
                i += 1
            else:
                buf = [s[2:]]
                i += 1
                while i < n and not lines[i].strip().endswith("$$"):
                    buf.append(lines[i])
                    i += 1
                if i < n:
                    buf.append(lines[i].strip()[:-2])
                    i += 1
                blocks.append(("math", " ".join(x.strip() for x in buf)))
        elif s.startswith("!["):
            m = re.match(r"!\[(.*?)\]\((.*?)\)", s)
            if m:
                blocks.append(("image", m.group(1), m.group(2)))
            i += 1
        elif s.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            blocks.append(("code", "\n".join(buf)))
        elif s.startswith("> "):
            buf = []
            while i < n and lines[i].strip().startswith("> "):
                buf.append(lines[i].strip()[2:])
                i += 1
            blocks.append(("quote", smart_join(buf)))
        elif re.match(r"^[-*]\s+", s):
            items = []
            while i < n and re.match(r"^[-*]\s+", lines[i].strip()):
                items.append(re.sub(r"^[-*]\s+", "", lines[i].strip()))
                i += 1
            blocks.append(("ul", items))
        elif re.match(r"^\d+\.\s+", s):
            items = []
            while i < n and re.match(r"^\d+\.\s+", lines[i].strip()):
                items.append(re.sub(r"^\d+\.\s+", "", lines[i].strip()))
                i += 1
            blocks.append(("ol", items))
        elif s.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.match(r"^:?-{2,}:?$", c.replace(" ", "")) for c in row if c):
                    rows.append(row)
                i += 1
            blocks.append(("table", rows))
        else:
            buf = [s]
            i += 1
            while i < n:
                t = lines[i].strip()
                if not t or t.startswith(SPECIAL_START) or re.match(r"^[-*]\s+", t) or re.match(r"^\d+\.\s+", t):
                    break
                buf.append(t)
                i += 1
            blocks.append(("p", smart_join(buf)))
    return blocks


# ----------------------------------------------------------------------------
# rendering
# ----------------------------------------------------------------------------
def add_text_runs(p, text, size=BODY_PT, color=INK, bold=False, italic=False,
                  ascii_font="Times New Roman", ea_font="宋体"):
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = p.add_run(part[1:-1])
            set_run_font(r, "Consolas", "宋体", size=size - 0.5, color=RGBColor(0x33, 0x41, 0x55))
            rPr = r._element.get_or_add_rPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "F1F5F9")
            rPr.append(shd)
        else:
            r = p.add_run(part)
            set_run_font(r, ascii_font, ea_font, size=size, bold=bold, italic=italic, color=color)


def add_math_runs(p, text, size=BODY_PT, color=INK, bold=False):
    for chunk, mode in math_segments(text):
        for piece, ital in split_italic(chunk):
            if not piece:
                continue
            r = p.add_run(piece)
            set_run_font(r, "Cambria Math", "宋体", size=size, bold=bold, italic=ital, color=color)
            if mode == "sub":
                r.font.subscript = True
            elif mode == "sup":
                r.font.superscript = True


def add_rich(p, text, size=BODY_PT, color=INK, bold=False):
    # split bold by **
    segs = re.split(r"(\*\*)", text)
    is_bold = bold
    for seg in segs:
        if seg == "**":
            is_bold = not is_bold
            continue
        if not seg:
            continue
        parts = re.split(r"(\$[^$]*\$)", seg)
        for part in parts:
            if not part:
                continue
            if part.startswith("$") and part.endswith("$") and len(part) > 2:
                add_math_runs(p, part[1:-1], size=size, color=color, bold=is_bold)
            else:
                add_text_runs(p, part, size=size, color=color, bold=is_bold)


def para(container, text, size=BODY_PT, indent=True, space_before=0, space_after=3,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, color=INK, bold=False, line=1.45):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    pf.alignment = align
    if indent:
        pf.first_line_indent = Pt(size * 2)
    add_rich(p, text, size=size, color=color, bold=bold)
    return p


def heading(doc, text, level):
    unnumbered = level == 2 and not re.match(r"^\d", text)
    if unnumbered:
        h = doc.add_paragraph(style="SectionTitle")
    else:
        h = doc.add_heading(level=level)
    add_rich(h, text, size={1: 21, 2: 15, 3: 11.5}[level], color=INK, bold=True)
    for r in h.runs:
        set_run_font(r, "Arial", "黑体", size={1: 21, 2: 15, 3: 11.5}[level], bold=True, color=INK)
    pf = h.paragraph_format
    pf.keep_with_next = True
    if level == 1:
        pf.page_break_before = True
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.space_before = Pt(0)
        pf.space_after = Pt(18)
        border_paragraph(h, color="1D4ED8", sz=8, side="bottom", space=6)
    elif level == 2:
        pf.space_before = Pt(16)
        pf.space_after = Pt(6)
    else:
        pf.space_before = Pt(10)
        pf.space_after = Pt(4)
    return h


def display_math(container, text, size=11):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(4)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.2
    add_math_runs(p, text, size=size)
    return p


def add_list(container, items, ordered=False, size=BODY_PT):
    for k, item in enumerate(items, 1):
        p = container.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Pt(size * 2.2)
        pf.first_line_indent = Pt(-size * 1.2)
        pf.space_before = Pt(1)
        pf.space_after = Pt(1)
        pf.line_spacing = 1.4
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        marker = f"{k}. " if ordered else "• "
        r = p.add_run(marker)
        set_run_font(r, "Times New Roman", "宋体", size=size, color=BLUE, bold=True)
        add_rich(p, item, size=size)


def add_quote(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.8)
    pf.right_indent = Cm(0.4)
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(6)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.4
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    shade_paragraph(p, "F8FAFC")
    border_paragraph(p, color="1D4ED8", sz=18, side="left", space=10)
    add_rich(p, text, size=9.8, color=RGBColor(0x33, 0x41, 0x55))


def add_code(doc, text):
    for line in text.split("\n"):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.6)
        pf.first_line_indent = Pt(0)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.15
        shade_paragraph(p, "F1F5F9")
        r = p.add_run(line if line else " ")
        set_run_font(r, "Consolas", "宋体", size=9, color=RGBColor(0x1E, 0x29, 0x3B))


def add_image(doc, caption, relpath):
    path = os.path.join(ROOT, relpath.replace("/", os.sep))
    if not os.path.exists(path):
        para(doc, f"[缺图：{relpath}]", size=9, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, color=GRAY)
        return
    from PIL import Image
    w, h = Image.open(path).size
    width = min(Cm(15.5), Cm(w / 110.0))
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.first_line_indent = Pt(0)
    p.add_run().add_picture(path, width=width)
    cap = doc.add_paragraph()
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(8)
    cap.paragraph_format.first_line_indent = Pt(0)
    add_rich(cap, caption, size=9, color=GRAY)


def add_table(doc, rows):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(len(rows), ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    def disp_len(s):
        s = (s or "").replace("$", "")
        s = re.sub(r"\\[A-Za-z]+", "", s)
        return sum(2 if is_cjk(ch) else 1 for ch in s)

    weights = []
    for j in range(ncols):
        w = max([disp_len(r[j]) for r in rows if j < len(r)] + [4])
        weights.append(max(6, w))
    total = sum(weights)
    width_cm = [16.2 * w / total for w in weights]
    widths = [Cm(x) for x in width_cm]
    for j in range(ncols):
        table.columns[j].width = widths[j]
    for i, row in enumerate(rows):
        row_cant_split(table.rows[i])
        for j in range(ncols):
            cell = table.rows[i].cells[j]
            cell.width = widths[j]
            # tight cell margins so tables fit on one page
            tcPr = cell._tc.get_or_add_tcPr()
            mar = OxmlElement("w:tcMar")
            for side, val in (("top", 30), ("left", 80), ("bottom", 30), ("right", 60)):
                el = OxmlElement("w:" + side)
                el.set(qn("w:w"), str(val))
                el.set(qn("w:type"), "dxa")
                mar.append(el)
            tcPr.append(mar)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            txt = row[j] if j < len(row) else ""
            fsize = 8.5 if len(rows) > 5 else 9
            if i == 0:
                add_rich(p, txt, size=fsize, bold=True)
                cell_shade(cell, "E2E8F0")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                add_rich(p, txt, size=fsize)
        if i == 0:
            trPr = table.rows[0]._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:tblHeader"))
    # borders
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "CBD5E1")
        borders.append(el)
    tblPr.append(borders)
    # spacing after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(0)
    sp.paragraph_format.line_spacing = 0.6


def add_box(doc, kind, title, blocks):
    fill, color = BOX.get(kind, ("F8FAFC", "64748B"))
    labels = {
        "definition": "定义", "theorem": "定理", "lemma": "引理",
        "corollary": "推论", "proposition": "命题", "example": "例",
        "remark": "注", "note": "注", "proof": "证明",
        "algorithm": "算法", "reading": "阅读", "history": "历史",
    }
    if title and kind in labels and not title.startswith(labels[kind]):
        title = labels[kind] + " " + title
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = CONTENT_W
    row = table.rows[0]
    row_cant_split(row)
    cell = row.cells[0]
    cell.width = CONTENT_W
    if fill:
        cell_shade(cell, fill)
    cell_borders(cell, color=color, sz=20 if kind != "proof" else 12, left_only=(kind != "proof"))
    cell_margins(cell, top=110, bottom=110, left=200, right=160)

    first = cell.paragraphs[0]
    first.paragraph_format.space_before = Pt(0)
    first.paragraph_format.space_after = Pt(3)
    first.paragraph_format.first_line_indent = Pt(0)
    first.paragraph_format.line_spacing = 1.25
    if title:
        add_rich(first, title, size=10.2, color=RGBColor.from_string(color), bold=True)
    else:
        first.paragraph_format.space_after = Pt(0)

    rendered_any = False
    last_text = ""
    for b in blocks:
        t = b[0]
        if t == "p":
            para(cell, b[1], size=10, space_after=3, line=1.4)
            rendered_any = True
            last_text = b[1]
        elif t == "math":
            display_math(cell, b[1], size=10.5)
            rendered_any = True
        elif t == "ul":
            add_list(cell, b[1], size=10)
            rendered_any = True
            last_text = b[1][-1] if b[1] else ""
        elif t == "ol":
            add_list(cell, b[1], ordered=True, size=10)
            rendered_any = True
            last_text = b[1][-1] if b[1] else ""
        elif t == "quote":
            p = cell.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            add_rich(p, b[1], size=9.5, color=GRAY)
            rendered_any = True
        elif t == "table":
            add_table(cell, b[1])
            rendered_any = True
        elif t == "code":
            for line in b[1].split("\n"):
                p = cell.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.3)
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.1
                r = p.add_run(line if line else " ")
                set_run_font(r, "Consolas", "宋体", size=8.5)
            rendered_any = True
    if kind == "proof" and "∎" not in last_text:
        p = cell.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run("∎")
        set_run_font(r, "Cambria Math", "宋体", size=10)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(0)
    sp.paragraph_format.line_spacing = 0.5


def render_blocks(doc, blocks):
    for b in blocks:
        t = b[0]
        if t == "h1":
            heading(doc, b[1], 1)
        elif t == "h2":
            heading(doc, b[1], 2)
        elif t == "h3":
            heading(doc, b[1], 3)
        elif t == "p":
            para(doc, b[1])
        elif t == "math":
            display_math(doc, b[1])
        elif t == "image":
            add_image(doc, b[1], b[2])
        elif t == "code":
            add_code(doc, b[1])
        elif t == "quote":
            add_quote(doc, b[1])
        elif t == "ul":
            add_list(doc, b[1])
        elif t == "ol":
            add_list(doc, b[1], ordered=True)
        elif t == "table":
            add_table(doc, b[1])
        elif t == "box":
            add_box(doc, b[1], b[2], b[3])


# ----------------------------------------------------------------------------
# document scaffolding
# ----------------------------------------------------------------------------
def setup_styles(doc):
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(BODY_PT)
    st.font.color.rgb = INK
    rPr = st.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), "宋体")
    pf = st.paragraph_format
    pf.line_spacing = 1.45
    pf.space_after = Pt(4)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    from docx.enum.style import WD_STYLE_TYPE
    names = [s.name for s in doc.styles]
    if "SectionTitle" not in names:
        st2 = doc.styles.add_style("SectionTitle", WD_STYLE_TYPE.PARAGRAPH)
        st2.base_style = doc.styles["Normal"]
        st2.font.name = "Arial"
        st2.font.size = Pt(15)
        st2.font.bold = True
        st2.font.color.rgb = INK
        rPr2 = st2.element.get_or_add_rPr()
        rFonts2 = rPr2.find(qn("w:rFonts"))
        if rFonts2 is None:
            rFonts2 = OxmlElement("w:rFonts")
            rPr2.append(rFonts2)
        rFonts2.set(qn("w:ascii"), "Arial")
        rFonts2.set(qn("w:hAnsi"), "Arial")
        rFonts2.set(qn("w:eastAsia"), "黑体")
        pf2 = st2.paragraph_format
        pf2.space_before = Pt(16)
        pf2.space_after = Pt(6)
        pf2.keep_with_next = True
        pf2.line_spacing = 1.2


def setup_section(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)
    sec.header_distance = Cm(1.2)
    sec.footer_distance = Cm(1.2)
    sec.different_first_page_header_footer = True

    # header
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run("AFA 与进程代数")
    set_run_font(r, "Arial", "黑体", size=8.5, color=GRAY)
    border_paragraph(hp, color="E2E8F0", sz=4, side="bottom", space=2)

    # footer page number
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("第 ")
    set_run_font(r, "Times New Roman", "宋体", size=9, color=GRAY)
    add_field(fp, "PAGE", "1")
    r = fp.add_run(" 页")
    set_run_font(r, "Times New Roman", "宋体", size=9, color=GRAY)


def title_page(doc, meta):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("A F A　与　进 程 代 数")
    set_run_font(r, "Arial", "黑体", size=11, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("AFA 与进程代数")
    set_run_font(r, "Arial", "黑体", size=36, color=INK, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("从反基础公理到超集合语义")
    set_run_font(r, "Arial", "黑体", size=17, color=GRAY)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(18)
    border_paragraph(p, color="1D4ED8", sz=24, side="bottom", space=1)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("“每个图都有唯一的装饰。”")
    set_run_font(r, "Times New Roman", "楷体", size=13, color=RGBColor(0x33, 0x41, 0x55))

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(30)
    r = p.add_run("—— Peter Aczel, Non-Well-Founded Sets, 1988")
    set_run_font(r, "Times New Roman", "宋体", size=10, color=GRAY)

    cover = os.path.join(ROOT, "figures", "fig-cover.png")
    if os.path.exists(cover):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(40)
        p.add_run().add_picture(cover, width=Cm(14.5))

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("自 学 教 材")
    set_run_font(r, "Arial", "黑体", size=12, color=INK, bold=True)

    p = doc.add_paragraph()
    r = p.add_run("2026")
    set_run_font(r, "Arial", "黑体", size=11, color=GRAY)


def toc_page(doc):
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("目　　录")
    set_run_font(r, "Arial", "黑体", size=18, color=INK, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    add_field(p, 'TOC \\o "1-2" \\h \\z \\u', "（在 Word 中右键 → 更新域，生成目录）")


def main(test=False):
    meta_path = os.path.join(os.path.dirname(__file__), "manifest.json")
    meta = json.load(open(meta_path, encoding="utf-8"))
    doc = Document()
    setup_styles(doc)
    setup_section(doc)
    title_page(doc, meta)
    toc_page(doc)

    files = [meta["preface"]] + meta["chapters"] + meta.get("appendices", [])
    if test:
        files = files[:8]
    for rel in files:
        path = os.path.join(ROOT, rel)
        if not os.path.exists(path):
            print("missing", path)
            continue
        lines = open(path, encoding="utf-8").read().splitlines()
        blocks = parse_blocks(lines)
        render_blocks(doc, blocks)
        print("rendered", rel, len(blocks), "blocks")

    out = os.path.join(ROOT, "out", "AFA与进程代数-教材.docx")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    doc.save(out)
    print("saved", out)


if __name__ == "__main__":
    main(test="--test" in sys.argv)
