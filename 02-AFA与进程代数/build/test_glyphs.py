# -*- coding: utf-8 -*-
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
samples = [
    "U+27E8:  ⟨a⟩⟨b⟩",
    "U+2329:  〈a〉〈b〉",
    "U+3008:  〈a〉〈b〉",
    "math:    P ∼ Q   ⊤   ∅   ℘   →",
    "ascii:   P ~ Q",
]
for s in samples:
    p = doc.add_paragraph()
    r = p.add_run(s)
    r.font.name = "Cambria Math"
    r.font.size = Pt(18)
    rPr = r._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rPr.append(rf)
    rf.set(qn("w:eastAsia"), "宋体")
out = r"D:\AI\AFARules\out\test-glyphs.docx"
doc.save(out)
print(out)
